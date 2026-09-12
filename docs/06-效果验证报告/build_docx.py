# -*- coding: utf-8 -*-
"""
由 06—效果验证报告.md 生成正式竞赛文档 06—效果验证报告.docx
================================================================
运行方式（需安装 python-docx，交付包自带 conda 环境已具备）：
    D:\\360Downloads\\anaconda\\envs\\programmind\\python.exe build_docx.py

支持的 Markdown 子集与自定义标记：
    # / ## / ###        标题
    | a | b |           表格（第二行为分隔行）
    - / 1.              列表
    **加粗**            行内加粗
    ![说明](路径)       插入图片并生成图注
    [[PAGEBREAK]]       分页
    [[CARDS]]           指标卡（其后若干行格式为：数值 | 说明，每行一张卡）
    [[FLOW]] 文本       流程图条（箭头串联）
    [[SHOT]] 文本       截图占位框
    > 文本              提示条
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(BASE_DIR, "06—效果验证报告.md")
DOCX_PATH = os.path.join(BASE_DIR, "06—效果验证报告.docx")

BRAND = RGBColor(0x16, 0x5D, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x6B, 0x72, 0x80)
HEAD_FILL = "165DFF"
CARD_FILL = "EEF4FF"
FLOW_FILL = "F5F8FF"
SHOT_FILL = "FAFBFC"
HEADER_TEXT = "ProgramMind · 高校计算机学科 AI-Native 教学科研平台"

COVER = {
    "title": "06—效果验证报告",
    "project": "ProgramMind",
    "slogan": "高校计算机学科 AI-Native 教学科研平台",
    "meta": [
        ("报告名称", "《06—效果验证报告》"),
        ("项目名称", "ProgramMind"),
        ("项目定位", "高校计算机学科 AI-Native 教学科研平台"),
        ("报告类型", "效果验证报告（竞赛 / 作品展示材料）"),
        ("验证日期", "2026 年 9 月 10 日"),
        ("验证方式", "真实环境部署运行 + 端到端业务链路实测 + 数据与日志留证"),
    ],
}

TOC = [
    "1. 验证概述",
    "2. 验证环境与部署说明",
    "3. 系统运行验证",
    "4. 核心功能验证",
    "5. 学生端典型任务验证",
    "6. 教师端典型任务验证",
    "7. 教学应用效果验证",
    "8. 系统性能与效率验证",
    "9. AI 交互效果验证",
    "10. 真实用户使用验证",
    "11. 综合效果评价",
    "12. 结论",
    "附录 A. 验证证据清单",
    "附录 B. 报告插图索引",
]


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------
def set_run_font(run, name_latin="Microsoft YaHei", name_east="Microsoft YaHei",
                 size=None, bold=None, color=None, italic=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name_latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name_latin)
    rfonts.set(qn("w:hAnsi"), name_latin)
    rfonts.set(qn("w:eastAsia"), name_east)


def shade(element, fill):
    """给单元格或段落加底色。"""
    pr = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def cell_shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D6DEF0", sz=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_field(paragraph, instr):
    """插入 Word 域，用于页码/总页数。"""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=GRAY)
    return run


def add_inline(paragraph, text, size=10.5, color=INK, bold=False,
               name_latin="Times New Roman", name_east="宋体"):
    """处理 **加粗** 行内标记。"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, name_latin, name_east, size=size, bold=True, color=BRAND)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, name_latin, name_east, size=size, bold=bold, color=color)


def paragraph(doc, text="", size=10.5, color=INK, bold=False, align=None,
              space_before=0, space_after=6, line_spacing=1.5,
              name_latin="Times New Roman", name_east="宋体", indent_first=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if align is not None:
        pf.alignment = align
    if indent_first:
        pf.first_line_indent = Pt(size * 2)
    if text:
        add_inline(p, text, size=size, color=color, bold=bold,
                   name_latin=name_latin, name_east=name_east)
    return p


def heading(doc, text, level):
    sizes = {1: 17, 2: 13.5, 3: 11.5}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 10)
    pf.space_after = Pt(8 if level == 1 else 6)
    pf.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei",
                 size=sizes[level], bold=True,
                 color=BRAND if level == 1 else INK)
    if level == 1:
        # 标题下加一条品牌色下边框
        ppr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), HEAD_FILL)
        borders.append(bottom)
        ppr.append(borders)
    return p


def page_break(doc):
    """
    插入分页符，但若上一段已经是分页符段落则跳过，
    避免 Markdown 中的 [[PAGEBREAK]] 与一级标题自动分页叠加产生空白页。
    """
    last_p = None
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            last_p = child
    if last_p is not None:
        xml = last_p.xml
        if 'w:type="page"' in xml and not "".join(last_p.itertext()).strip():
            return
    doc.add_page_break()


# --------------------------------------------------------------------------
# 封面 / 页眉页脚 / 目录
# --------------------------------------------------------------------------
def build_cover(doc):
    for _ in range(3):
        paragraph(doc, "", space_after=0)
    p = paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    run = p.add_run(COVER["project"])
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=30, bold=True, color=BRAND)
    p = paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    run = p.add_run(COVER["slogan"])
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=13, color=GRAY)
    p = paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=24)
    run = p.add_run(COVER["title"])
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=26, bold=True, color=INK)

    table = doc.add_table(rows=len(COVER["meta"]), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (k, v) in enumerate(COVER["meta"]):
        row = table.rows[i]
        for j, cell in enumerate(row.cells):
            cell.width = Cm(3.6 if j == 0 else 10.4)
            set_cell_borders(cell, color="FFFFFF")
        c0, c1 = row.cells
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        set_run_font(r0, "Microsoft YaHei", "Microsoft YaHei", size=10.5, bold=True, color=GRAY)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        set_run_font(r1, "Microsoft YaHei", "Microsoft YaHei", size=10.5, color=INK)
    for _ in range(4):
        paragraph(doc, "", space_after=0)
    paragraph(doc, "ProgramMind 项目组", align=WD_ALIGN_PARAGRAPH.CENTER,
              size=11, color=GRAY, space_after=0)
    paragraph(doc, "2026 年 9 月", align=WD_ALIGN_PARAGRAPH.CENTER,
              size=11, color=GRAY)
    page_break(doc)


def build_toc(doc):
    heading(doc, "目录", 1)
    for item in TOC:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.line_spacing = 1.4
        run = p.add_run(item)
        set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=10.5, color=INK)
    page_break(doc)


def build_header_footer(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(HEADER_TEXT)
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=8.5, color=GRAY)
    ppr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "D6DEF0")
    borders.append(bottom)
    ppr.append(borders)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("第 ")
    set_run_font(r1, "Microsoft YaHei", "Microsoft YaHei", size=9, color=GRAY)
    add_field(fp, "PAGE")
    r2 = fp.add_run(" 页 / 共 ")
    set_run_font(r2, "Microsoft YaHei", "Microsoft YaHei", size=9, color=GRAY)
    add_field(fp, "NUMPAGES")
    r3 = fp.add_run(" 页")
    set_run_font(r3, "Microsoft YaHei", "Microsoft YaHei", size=9, color=GRAY)


# --------------------------------------------------------------------------
# 复合元素
# --------------------------------------------------------------------------
def add_cards(doc, items):
    per_row = 4
    rows = (len(items) + per_row - 1) // per_row
    table = doc.add_table(rows=rows * 2, cols=per_row)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (value, label) in enumerate(items):
        r, c = divmod(idx, per_row)
        v_cell = table.cell(r * 2, c)
        l_cell = table.cell(r * 2 + 1, c)
        for cell in (v_cell, l_cell):
            set_cell_borders(cell, color="D6DEF0", sz=6)
            cell_shade(cell, CARD_FILL)
        vp = v_cell.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vr = vp.add_run(value)
        set_run_font(vr, "Microsoft YaHei", "Microsoft YaHei", size=15, bold=True, color=BRAND)
        lp = l_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lr = lp.add_run(label)
        set_run_font(lr, "Microsoft YaHei", "Microsoft YaHei", size=9, color=GRAY)
    paragraph(doc, "", size=6, space_after=2)


def add_flow(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_borders(cell, color="BEDAFF", sz=8)
    cell_shade(cell, FLOW_FILL)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=10.5, bold=True, color=BRAND)
    paragraph(doc, "", size=6, space_after=2)


def add_shot(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_borders(cell, color="C9D4E8", sz=8)
    cell_shade(cell, SHOT_FILL)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(text)
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=10, color=GRAY, italic=True)
    paragraph(doc, "", size=6, space_after=2)


def add_image(doc, caption, path):
    full = os.path.join(BASE_DIR, path.replace("/", os.sep))
    if not os.path.exists(full):
        add_shot(doc, "【缺失图片】" + caption + "（" + path + "）")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(full, width=Cm(15.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    set_run_font(r, "Microsoft YaHei", "Microsoft YaHei", size=9, color=GRAY)


def add_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell_shade(cell, HEAD_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=10, bold=True,
                     color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx in range(len(header)):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = cells[c_idx]
            if r_idx % 2 == 1:
                cell_shade(cell, "F7FAFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, text, size=9.5)
    paragraph(doc, "", size=6, space_after=4)


def add_bullet(doc, text, numbered=False, level=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(14 + level * 12)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.4
    marker = "" if numbered else "• "
    run = p.add_run(marker)
    set_run_font(run, "Microsoft YaHei", "Microsoft YaHei", size=10.5, color=BRAND)
    add_inline(p, text, size=10.5)


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_borders(cell, color="BEDAFF", sz=8)
    cell_shade(cell, FLOW_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, text, size=9.5, color=INK)
    paragraph(doc, "", size=6, space_after=2)


# --------------------------------------------------------------------------
# 主流程
# --------------------------------------------------------------------------
def parse_cards(lines, i):
    """读取 [[CARDS]] 之后的 数值|说明 行。"""
    items = []
    j = i + 1
    while j < len(lines):
        line = lines[j].strip()
        if not line or "|" not in line:
            break
        value, label = line.split("|", 1)
        items.append((value.strip(), label.strip()))
        j += 1
    return items, j


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        rows.append([c.strip() for c in raw.split("|")])
        i += 1
    if len(rows) >= 2 and set("".join(rows[1])) <= set("-: "):
        rows = [rows[0]] + rows[2:]
    return rows, i


def build_body(doc, lines):
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].rstrip("\n")
        s = line.strip()

        if not s or s.startswith("<!--"):
            i += 1
            continue

        if s == "[[PAGEBREAK]]":
            page_break(doc)
            i += 1
            continue

        if s.startswith("[[CARDS]]"):
            items, i = parse_cards(lines, i)
            add_cards(doc, items)
            continue

        if s.startswith("[[FLOW]]"):
            add_flow(doc, s[len("[[FLOW]]"):].strip())
            i += 1
            continue

        if s.startswith("[[SHOT]]"):
            add_shot(doc, s[len("[[SHOT]]"):].strip())
            i += 1
            continue

        m = re.match(r"^!\[(.*?)\]\((.+?)\)$", s)
        if m:
            add_image(doc, m.group(1), m.group(2))
            i += 1
            continue

        if s.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows[0], rows[1:])
            continue

        if s.startswith("### "):
            heading(doc, s[4:], 3)
            i += 1
            continue
        if s.startswith("## "):
            heading(doc, s[3:], 2)
            i += 1
            continue
        if s.startswith("# "):
            if first_heading:
                first_heading = False
            else:
                page_break(doc)
            heading(doc, s[2:], 1)
            i += 1
            continue

        if s.startswith("> "):
            add_note(doc, s[2:])
            i += 1
            continue

        if s in ("---", "***"):
            i += 1
            continue

        if re.match(r"^\d+\.\s", s):
            add_bullet(doc, re.sub(r"^\d+\.\s", "", s), numbered=False)
            i += 1
            continue

        if s.startswith("- "):
            add_bullet(doc, s[2:])
            i += 1
            continue

        paragraph(doc, s, indent_first=False)
        i += 1


def main():
    if not os.path.exists(MD_PATH):
        print("找不到 Markdown 源文件：" + MD_PATH)
        sys.exit(1)
    with open(MD_PATH, "r", encoding="utf-8") as f:
        text = f.read()
    lines = text.splitlines()
    body_start = 0
    for idx, ln in enumerate(lines):
        if ln.strip() == "<!-- BODY -->":
            body_start = idx + 1
            break

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    build_header_footer(doc)
    build_cover(doc)
    build_toc(doc)
    build_body(doc, lines[body_start:])

    doc.save(DOCX_PATH)
    print("已生成：" + DOCX_PATH)
    print("大小：%.1f KB" % (os.path.getsize(DOCX_PATH) / 1024.0))


if __name__ == "__main__":
    main()
